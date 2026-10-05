"""Excel analysis reports (HANDOVER §12.1, §12.2). Does file I/O.

One question: Özet · Öneriler · Alan Detayları · Notlar ve Öneriler.
A term list (ADR-033): the same four sheets, combined, with a term column.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from vsa.models import (
    AnalysisResult,
    DictColumn,
    FlagKind,
    Level,
    ListItem,
    ListResult,
    ObjectMatch,
    Verdict,
)
from vsa.text.normalize import fold

CellValue = str | int | float | None

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE_FONT = Font(name=FONT, size=14, bold=True, color="1F3864")
SECTION_FONT = Font(name=FONT, size=11, bold=True, color="1F3864")
SMALL_FONT = Font(name=FONT, size=9, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

LEVEL_FILL = {
    Level.HIGH: PatternFill("solid", fgColor="C6EFCE"),
    Level.MEDIUM: PatternFill("solid", fgColor="FFEB9C"),
    Level.LOW: PatternFill("solid", fgColor="FFC7CE"),
}
VERDICT_FILL = {
    Verdict.FOUND: PatternFill("solid", fgColor="C6EFCE"),
    Verdict.PARTIAL: PatternFill("solid", fgColor="FFEB9C"),
    Verdict.NOT_FOUND: PatternFill("solid", fgColor="FFC7CE"),
}

FLAG_LABEL = {
    FlagKind.MODEL_ESTIMATED: "Model tahmini — doğrulanmalı",
    FlagKind.CORRECTED: "Kaynak açıklama düzeltildi",
    FlagKind.NEEDS_VERIFICATION: "Doğrulanmalı",
    FlagKind.NAMING_MISMATCH: "İsim/içerik uyumsuzluğu",
    FlagKind.QUALITY_NOTE: "Kalite notu",
}


def slugify(text: str, limit: int = 40) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", fold(text[:limit])).strip("_")
    return slug or "talep"


def report_path(out_dir: Path, mode: str, query: str, now: datetime | None = None) -> Path:
    stamp = (now or datetime.now()).strftime("%Y%m%d_%H%M")
    return out_dir / f"VSA_{mode}_{slugify(query)}_{stamp}.xlsx"


def _table(
    ws: Worksheet,
    start_row: int,
    headers: Sequence[str],
    rows: Sequence[Sequence[CellValue]],
    widths: Sequence[int] | None = None,
) -> int:
    """Write a styled table; returns the next free row."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(rows, start_row + 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font, cell.border, cell.alignment = BODY_FONT, BORDER, WRAP
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    return start_row + len(rows) + 1


def _finish_sheet(ws: Worksheet, header_row: int, n_cols: int, n_rows: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{header_row + 1}"
    if n_rows:
        last = get_column_letter(n_cols)
        ws.auto_filter.ref = f"A{header_row}:{last}{header_row + n_rows}"


def _merged(ws: Worksheet, row: int, first: int, last: int, value: CellValue) -> Cell:
    """Value across merged cells ``first..last`` of a row, wrapped."""
    cell = ws.cell(row=row, column=first, value=value)
    cell.font, cell.alignment = BODY_FONT, WRAP
    if last > first:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    return cell


def _grow(ws: Worksheet, row: int, text: str, chars_per_line: int) -> None:
    """Merged cells do not auto-size in Excel: give long texts room."""
    lines = sum(max(1, -(-len(part) // chars_per_line)) for part in text.split("\n"))
    ws.row_dimensions[row].height = max(15, 14 * lines)


def _section(ws: Worksheet, row: int, title: str) -> int:
    ws.cell(row=row, column=1, value=title).font = SECTION_FONT
    return row + 1


def _bullets(ws: Worksheet, row: int, title: str, items: Sequence[str]) -> int:
    if not items:
        return row
    row = _section(ws, row, title)
    for item in items:
        _merged(ws, row, 1, 5, f"• {item}")
        _grow(ws, row, item, 150)
        row += 1
    return row + 1


def _covers(m: ObjectMatch) -> str:
    if m.covers:
        return m.covers
    return ", ".join(m.covered) or ", ".join(h.col.column for h in m.columns[:4])


WARN_FILL = PatternFill("solid", fgColor="FFEB9C")


def answered_by(r: AnalysisResult) -> str:
    """Who wrote the answer, said plainly (ADR-034)."""
    if r.analyst:
        return f"Analist (dil modeli: {r.llm_model})"
    if r.fallback:
        return f"Analiz yapılamadı — {r.fallback}"
    cut = f"güven skorları %{round((1 - r.confidence_factor) * 100)} düşürülerek gösteriliyor"
    return f"Kural motoru (yalnız ölçüm amaçlı; kullanıcıya gösterilmez) — {cut}."


def _dictionary_line(r: AnalysisResult) -> str:
    rows = r.dictionary_version.rsplit("-", 1)[-1]
    size = f"{int(rows):,}".replace(",", ".") + " satır" if rows.isdigit() else r.dictionary_version
    objects = f", {r.dictionary_objects} obje" if r.dictionary_objects else ""
    return f"Kaynak sözlük: {r.dictionary_source} ({size}{objects})"


def _summary(ws: Worksheet, r: AnalysisResult) -> None:
    """Özet: request, verdict, suggestion table, design, traps, scale (HANDOVER §12.2)."""
    ws.title = "Özet"
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDE", (16, 46, 70, 10, 12), strict=True):
        ws.column_dimensions[col].width = width
    ws["A1"] = "Veri Ambarı Veri Sözlüğü – İş Birimi Talep Analizi"
    ws["A1"].font = TITLE_FONT
    _merged(ws, 2, 1, 3, _dictionary_line(r)).font = SMALL_FONT
    stamp = datetime.strptime(r.generated_at, "%Y-%m-%d %H:%M").strftime("%d.%m.%Y %H:%M")
    _merged(ws, 2, 4, 5, f"Üretim: {stamp}").font = SMALL_FONT

    row = 4
    info = [("Talep", r.query), ("Sonuç", r.summary), ("Cevabı üreten", answered_by(r))]
    if r.interpretation:
        info.append(("Talebin yorumu", r.interpretation))
    for label, value in info:
        ws.cell(row=row, column=1, value=label).font = BOLD
        ws.cell(row=row, column=1).alignment = WRAP
        cell = _merged(ws, row, 2, 5, value)
        _grow(ws, row, value, 120)
        if label == "Sonuç":
            cell.fill, cell.font = VERDICT_FILL[r.verdict], BOLD
        if label == "Cevabı üreten" and not r.analyst:
            cell.fill = WARN_FILL
        row += 1

    row = _section(ws, row + 1, "Öneri Özeti")
    rows: list[list[CellValue]] = [
        [i, m.object_key, _covers(m), round(m.score, 2), m.level.value]
        for i, m in enumerate(r.objects, 1)
    ]
    if not rows:
        rows = [["-", "Sözlükte karşılığı bulunamadı", "-", "-", "-"]]
    start = row
    headers = ["Sıra", "Veritabanı.Şema.Obje", "Kapsadığı Bilgi", "Güven", "Seviye"]
    row = _table(ws, start, headers, rows)
    for i in range(start + 1, row):
        ws.cell(row=i, column=4).number_format = "0.00"
        for lvl in Level:
            if ws.cell(row=i, column=5).value == lvl.value:
                ws.cell(row=i, column=5).fill = LEVEL_FILL[lvl]

    row = _bullets(ws, row + 1, "Önerilen Kurgu", r.design)
    row = _bullets(ws, row, "Dikkat Edilmesi Gerekenler", r.attention)

    row = _section(ws, row, "Güven Skoru Ölçeği")
    start = row
    row = _table(
        ws,
        start,
        ["Seviye", "Aralık", "Anlamı"],
        [
            ["Yüksek", "%80–100", "Alan adı veya açıklaması sorulan kavramla doğrudan örtüşüyor."],
            ["Orta", "%50–79", "Anlam olarak ilişkili; kapsam, granülerlik veya tanım farkı var."],
            ["Düşük", "%0–49", "Dolaylı ilişki var; sorulan veri bu alandan türetilebilir."],
        ],
    )
    for i, lvl in enumerate(Level):
        ws.cell(row=start + 1 + i, column=1).fill = LEVEL_FILL[lvl]

    row = _section(ws, row + 1, "Yöntem Notları")
    notes = [
        *r.method,
        "Talep kavramları: " + (", ".join(r.concepts) or "-"),
        f"Süre: {r.elapsed_ms / 1000:.1f} sn",
    ]
    for n in notes:
        _merged(ws, row, 1, 5, n).font = SMALL_FONT
        row += 1


def _suggestions(ws: Worksheet, r: AnalysisResult) -> None:
    headers = [
        "Sıra", "Veritabanı", "Şema", "Obje", "Veri Seti Grubu", "İlgili Alanlar",
        "Gerekçe", "Kısıt / Dikkat", "Güven Skoru", "Güven Seviyesi", "Kullanım Önerisi",
    ]  # fmt: skip
    ws["A1"] = f"Öneriler – Güven Skoruna Göre İlk {max(len(r.objects), 1)}"
    ws["A1"].font = TITLE_FONT
    ws["A2"], ws["B2"] = "Talep:", r.query
    ws["A2"].font, ws["B2"].font = BOLD, BODY_FONT
    rows: list[list[CellValue]] = [
        [
            i,
            m.database,
            m.schema,
            m.object_name,
            ", ".join(m.dataset_groups) or "-",
            "\n".join(h.col.column for h in m.columns),
            m.reason,
            m.caveat,
            round(m.score, 2),
            m.level.value,
            m.usage,
        ]
        for i, m in enumerate(r.objects, 1)
    ]
    if not rows:
        rows = [["-", "-", "-", "-", "-", "-", r.summary, "-", 0, "-", "-"]]
    header_row = 3
    _table(ws, header_row, headers, rows, [6, 10, 8, 30, 20, 30, 70, 60, 10, 11, 40])
    for i in range(header_row + 1, header_row + len(rows) + 1):
        ws.cell(row=i, column=9).number_format = "0.00"
        level = ws.cell(row=i, column=10).value
        for lvl in Level:
            if level == lvl.value:
                ws.cell(row=i, column=10).fill = LEVEL_FILL[lvl]
    _finish_sheet(ws, header_row, len(headers), len(rows))


def _column_text(col: DictColumn) -> str:
    text = col.description
    if col.synonyms:
        text += f" Eş anlamlılar/aranabilir terimler: {', '.join(col.synonyms)}."
    return text


def _field_details(ws: Worksheet, r: AnalysisResult) -> None:
    ws["A1"] = "Önerilen Alanların Sözlükteki Tanımları"
    ws["A1"].font = TITLE_FONT
    headers = ["Sıra", "Veritabanı.Şema.Obje", "Alan Adı", "Sözlük Açıklaması", "Kalite Bayrağı"]
    rows: list[list[CellValue]] = []
    for i, m in enumerate(r.objects, 1):
        for h in m.columns:
            flags = "\n".join(
                f"{FLAG_LABEL[f.kind]}: {f.text}"
                if f.kind is not FlagKind.MODEL_ESTIMATED
                else FLAG_LABEL[f.kind]
                for f in h.col.flags
            )
            rows.append([i, m.object_key, h.col.column, _column_text(h.col), flags or "-"])
    _table(ws, 2, headers, rows, [6, 42, 30, 100, 34])
    _finish_sheet(ws, 2, len(headers), len(rows))


def _notes(ws: Worksheet, r: AnalysisResult) -> None:
    ws["A1"] = "Notlar ve Öneriler"
    ws["A1"].font = TITLE_FONT
    headers = ["Kapsam", "Başlık", "Açıklama"]
    rows: list[list[CellValue]] = [[n.scope, n.title, n.text] for n in r.notes]
    _table(ws, 2, headers, rows, [16, 40, 110])
    _finish_sheet(ws, 2, len(headers), len(rows))


def write_ask_report(result: AnalysisResult, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    _summary(ws, result)
    _suggestions(wb.create_sheet("Öneriler"), result)
    _field_details(wb.create_sheet("Alan Detayları"), result)
    _notes(wb.create_sheet("Notlar ve Öneriler"), result)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --------------------------------------------------------------------------- list search (ADR-033)


def _level_fill(ws: Worksheet, col: int, first: int, last: int) -> None:
    for row in range(first, last + 1):
        cell = ws.cell(row=row, column=col)
        for lvl, fill in LEVEL_FILL.items():
            if cell.value == lvl.value:
                cell.fill = fill
        for verdict, vfill in VERDICT_FILL.items():
            if cell.value == verdict.value:
                cell.fill = vfill


def _item_state(item: ListItem, cancelled: bool) -> str:
    if item.result is not None:
        return item.result.verdict.value
    if item.error:
        return "HATA"
    return "DURDURULDU" if cancelled else "BEKLİYOR"


def _list_summary(ws: Worksheet, r: ListResult, first_rows: dict[int, int]) -> None:
    ws.title = "Özet"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Veri Ambarı Veri Sözlüğü – Toplu Talep Analizi"
    ws["A1"].font = TITLE_FONT
    done = [i.result for i in r.items if i.result is not None]
    counts = {v: sum(1 for d in done if d.verdict is v) for v in Verdict}
    stamp = datetime.strptime(r.generated_at, "%Y-%m-%d %H:%M").strftime("%d.%m.%Y %H:%M")
    outcome = (
        f"{len(r.items)} terim: {counts[Verdict.FOUND]} VAR, {counts[Verdict.PARTIAL]} KISMEN VAR, "
        f"{counts[Verdict.NOT_FOUND]} BULUNAMADI"
    )
    if len(done) < len(r.items):
        outcome += f", {len(r.items) - len(done)} cevapsız"
    if r.cancelled:
        outcome += " — liste durduruldu, rapor yarım"
    failed = [d for d in done if not d.analyst]
    method = f"{len(done) - len(failed)} terim analist (dil modeli)"
    if failed:
        method += (f", {len(failed)} terim analiz edilemedi — model bağlı değil ya da "
                   "çalışamadı (Yöntem sütunu)")  # fmt: skip
    lines = [
        ("Liste", r.name),
        ("Kaynak sözlük", f"{r.dictionary_source} (sürüm {r.dictionary_version})"),
        ("Üretim", stamp),
        ("Sonuç", outcome),
        ("Cevabı üreten", method),
    ]
    if r.notes:
        lines.append(("Dosya", " ".join(r.notes)))
    for row, (label, value) in enumerate(lines, 3):
        ws.cell(row=row, column=1, value=label).font = BOLD
        cell = _merged(ws, row, 2, 9, value)
        if label == "Cevabı üreten" and failed:
            cell.fill = WARN_FILL
    start = len(lines) + 4
    headers = [
        "#", "Terim", "Sonuç", "En İyi Tablo", "Güven", "Seviye", "Özet", "Yöntem", "Süre (sn)",
    ]  # fmt: skip
    rows: list[list[CellValue]] = []
    for item in r.items:
        res = item.result
        best = res.objects[0] if res and res.objects else None
        if res is None:
            method = "-"
        else:
            method = "Analist" if res.analyst else "Yapılamadı" if res.fallback else "Kural"
        rows.append([
            item.index, item.term, _item_state(item, r.cancelled),
            best.object_key if best else "-",
            round(best.score, 2) if best else None,
            best.level.value if best else "-",
            res.summary if res else item.error or "-",
            method, round(item.elapsed_ms / 1000, 1) if item.elapsed_ms else None,
        ])  # fmt: skip
    _table(ws, start, headers, rows, [5, 34, 14, 44, 8, 9, 80, 13, 9])
    for n, item in enumerate(r.items):
        row = start + 1 + n
        ws.cell(row=row, column=5).number_format = "0.00"
        target = first_rows.get(item.index)
        if target:  # the term jumps to its suggestions
            term = ws.cell(row=row, column=2)
            term.hyperlink = f"#'Öneriler'!A{target}"
            term.font = Font(name=FONT, size=10, color="1F3864", underline="single")
    _level_fill(ws, 3, start + 1, start + len(rows))
    _level_fill(ws, 6, start + 1, start + len(rows))
    _finish_sheet(ws, start, len(headers), len(rows))


def _list_suggestions(ws: Worksheet, r: ListResult) -> dict[int, int]:
    """Every term's suggestions one under the other; returns term index -> first row."""
    headers = [
        "#", "Terim", "Sıra", "Veritabanı.Şema.Obje", "Veri Seti Grubu", "Kapsadığı Bilgi",
        "İlgili Alanlar", "Gerekçe", "Kısıt / Dikkat", "Güven", "Seviye", "Kullanım Önerisi",
    ]  # fmt: skip
    rows: list[list[CellValue]] = []
    first: dict[int, int] = {}
    for item in r.items:
        first[item.index] = len(rows) + 2
        res = item.result
        if res is None or not res.objects:
            text = res.summary if res else item.error or _item_state(item, r.cancelled)
            rows.append([item.index, item.term, "-", "-", "-", "-", "-", text, "-", None, "-", "-"])
            continue
        for rank, m in enumerate(res.objects, 1):
            rows.append([
                item.index, item.term, rank, m.object_key, ", ".join(m.dataset_groups) or "-",
                _covers(m), "\n".join(h.col.column for h in m.columns), m.reason, m.caveat,
                round(m.score, 2), m.level.value, m.usage,
            ])  # fmt: skip
    _table(ws, 1, headers, rows, [5, 30, 6, 42, 20, 30, 30, 70, 60, 8, 9, 40])
    for i in range(2, len(rows) + 2):
        ws.cell(row=i, column=10).number_format = "0.00"
    _level_fill(ws, 11, 2, len(rows) + 1)
    _finish_sheet(ws, 1, len(headers), len(rows))
    return first


def _list_fields(ws: Worksheet, r: ListResult) -> None:
    """Suggested columns with their dictionary text, each once, with the terms it serves."""
    seen: dict[str, tuple[str, DictColumn, list[int]]] = {}
    for item in r.items:
        for m in item.result.objects if item.result else []:
            for h in m.columns:
                entry = seen.setdefault(h.col.key, (m.object_key, h.col, []))
                if item.index not in entry[2]:
                    entry[2].append(item.index)
    rows: list[list[CellValue]] = []
    for obj, col, terms in seen.values():
        flags = "\n".join(FLAG_LABEL[f.kind] for f in col.flags) or "-"
        rows.append([obj, col.column, _column_text(col), flags, ", ".join(map(str, terms))])
    headers = ["Veritabanı.Şema.Obje", "Alan Adı", "Sözlük Açıklaması", "Kalite Bayrağı", "Terim #"]
    _table(ws, 1, headers, rows, [42, 30, 100, 30, 10])
    _finish_sheet(ws, 1, len(headers), len(rows))


def _list_notes(ws: Worksheet, r: ListResult) -> None:
    rows: list[list[CellValue]] = []
    for item in r.items:
        res = item.result
        if res is None:
            continue
        rows += [[item.index, item.term, "Önerilen kurgu", "", d] for d in res.design]
        rows += [[item.index, item.term, "Dikkat", "", a] for a in res.attention]
        rows += [[item.index, item.term, n.scope, n.title, n.text] for n in res.notes]
    headers = ["#", "Terim", "Kapsam", "Başlık", "Açıklama"]
    _table(ws, 1, headers, rows, [5, 30, 16, 36, 100])
    _finish_sheet(ws, 1, len(headers), len(rows))


def write_list_report(result: ListResult, path: Path) -> Path:
    """One workbook for a whole term list: Özet (a row per term) · Öneriler · Alan
    Detayları · Notlar — combined sheets with a term column, filterable (ADR-033)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    first_rows = _list_suggestions(wb.create_sheet("Öneriler"), result)
    _list_summary(ws, result, first_rows)
    _list_fields(wb.create_sheet("Alan Detayları"), result)
    _list_notes(wb.create_sheet("Notlar ve Öneriler"), result)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
