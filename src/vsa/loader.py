"""Parsing the dictionary template, term dictionary and stopwords.

This is one of the few modules allowed to do I/O. Parsing helpers are pure and tested
on their own. Workbooks are read with openpyxl (read-only, values only); the app's
dictionary itself is ``dictionary_store`` (ADR-050), which imports workbooks with these.
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import re
from collections.abc import Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from vsa.models import DictColumn, ObjectProfile, TermGroup
from vsa.text.normalize import fold, load_stopwords

log = logging.getLogger(__name__)

_PII_MARKER = re.compile(r"kvkk|kişisel\s+veri", re.IGNORECASE)

# Template column names, matched case-insensitively.
_REQUIRED = {
    "database": ("databasename",),
    "schema": ("schemaname",),
    "object": ("objectname",),
    "column": ("columnname",),
    "description": ("columndescription",),
}
_OPTIONAL = {
    "synonyms": ("synonyms",),
    "role": ("role",),
    "summary": ("summary",),
}

Row = Mapping[str, object]


# --------------------------------------------------------------------------- parsing


def split_synonyms(text: str) -> list[str]:
    """Comma-split that ignores commas inside parentheses; trims a trailing period."""
    items: list[str] = []
    depth = 0
    buf: list[str] = []
    for ch in text:
        if ch in "([":
            depth += 1
        elif ch in ")]" and depth:
            depth -= 1
        if ch == "," and depth == 0:
            items.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    items.append("".join(buf))
    cleaned = [s.strip().rstrip(".").strip() for s in items]
    return [s for s in cleaned if s]


def _resolve_columns(headers: Iterable[str]) -> dict[str, str]:
    """Map logical names -> actual header names, case-insensitively."""
    headers = list(headers)
    lookup = {fold(str(c)).replace(" ", ""): str(c) for c in headers}
    resolved: dict[str, str] = {}
    for logical, aliases in {**_REQUIRED, **_OPTIONAL}.items():
        for alias in aliases:
            if alias in lookup:
                resolved[logical] = lookup[alias]
                break
        else:
            if logical in _REQUIRED:
                raise ValueError(
                    f"Sözlükte zorunlu kolon bulunamadı: {aliases[0]} "
                    f"(mevcut: {', '.join(map(str, headers))})"
                )
    return resolved


def _cell(value: object) -> str:
    if value is None or (isinstance(value, float) and value != value):  # empty / NaN
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # a number cell reads "5", not "5.0"
    return str(value).strip()


OBJECTS_SHEET = "Objeler"


def _names(cell: object) -> list[str]:
    return [n.strip() for n in _cell(cell).split(",") if n.strip()]


def object_profiles(objects: Iterable[Row]) -> dict[str, ObjectProfile]:
    """The object sheet as profiles keyed by ``DB.Schema.Object``. Pure."""
    out: dict[str, ObjectProfile] = {}
    for rec in objects:
        key = _cell(rec.get("ObjectKey")) or ".".join(
            _cell(rec.get(c)) for c in ("DatabaseName", "SchemaName", "ObjectName")
        )
        if not key.strip("."):
            continue
        out[key] = ObjectProfile(
            key=key,
            description=_cell(rec.get("ObjectDescription")),
            grain=_cell(rec.get("Grain")),
            time_columns=tuple(_names(rec.get("TimeColumns"))),
            domain=_cell(rec.get("BusinessDomain")),
            group=_cell(rec.get("DatasetGroup")),
        )
    return out


def build_columns(
    rows: Sequence[Row], groups: Mapping[str, str] | None = None
) -> tuple[list[DictColumn], list[str]]:
    """Turn dictionary rows into DictColumns. Pure. A column's DatasetGroup comes from
    its object (``groups``, read from the object sheet)."""
    names = _resolve_columns(rows[0].keys() if rows else ())
    groups = groups or {}
    warnings: list[str] = []

    columns: list[DictColumn] = []
    seen: set[str] = set()
    for i, rec in enumerate(rows):
        db, schema = _cell(rec.get(names["database"])), _cell(rec.get(names["schema"]))
        obj, col = _cell(rec.get(names["object"])), _cell(rec.get(names["column"]))
        description = _cell(rec.get(names["description"]))
        if not (db and schema and obj and col):
            warnings.append(f"Satır {i + 2}: eksik veritabanı/şema/obje/kolon, atlandı")
            continue
        key = f"{db}.{schema}.{obj}.{col}"
        if key in seen:
            warnings.append(f"Tekrarlanan kolon, ilk kayıt tutuldu: {key}")
            continue
        seen.add(key)

        listed = _cell(rec.get(names["synonyms"])) if "synonyms" in names else ""
        columns.append(
            DictColumn(
                id=len(columns),
                database=db,
                schema=schema,
                object_name=obj,
                column=col,
                description=description,
                synonyms=tuple(split_synonyms(listed)),
                dataset_group=groups.get(f"{db}.{schema}.{obj}") or None,
                has_pii=bool(_PII_MARKER.search(f"{description} {listed}")),
                role=_cell(rec.get(names["role"])) if "role" in names else "",
                summary=_cell(rec.get(names["summary"])) if "summary" in names else "",
            )
        )
    return columns, warnings


# --------------------------------------------------------------------------- I/O


def file_version(path: Path, row_count: int) -> str:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()[:12]
    return f"{digest}-{row_count}"


class TooManyRows(ValueError):
    pass


def _sheet_rows(
    sheet: Iterable[Sequence[object]], max_rows: int | None = None
) -> list[list[object]]:
    """Every row of a sheet as a list; trailing empty rows dropped. More than ``max_rows``
    rows raises ``TooManyRows`` before the rest is read (a crafted upload cannot make the
    server unpack millions of rows)."""
    rows: list[list[object]] = []
    for r in sheet:
        if max_rows is not None and len(rows) >= max_rows:
            if any(_cell(c) for c in r):
                raise TooManyRows(str(max_rows))
            continue
        rows.append(list(r))
    while rows and not any(_cell(c) for c in rows[-1]):
        rows.pop()
    return rows


def records(rows: Sequence[Sequence[object]]) -> Iterator[dict[str, object]]:
    """A header row and the rows below it, as dicts keyed by the header."""
    if not rows:
        return
    header = [_cell(h) for h in rows[0]]
    for row in rows[1:]:
        yield {h: v for h, v in zip(header, row, strict=False) if h}


def read_workbook(
    path: Path, max_rows: int | None = None, max_cols: int | None = None
) -> dict[str, list[list[object]]]:
    """Every sheet of a workbook as rows of cell values (at most ``max_cols`` wide)."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            ws.title: _sheet_rows(ws.iter_rows(max_col=max_cols, values_only=True), max_rows)
            for ws in book.worksheets
        }
    finally:
        book.close()


def _jsonl(path: Path) -> Iterator[dict[str, object]]:
    """The records of a jsonl file; blank lines are skipped, a broken line is an error."""
    with path.open(encoding="utf-8-sig") as fh:
        for n, line in enumerate(fh, 1):
            if not line.strip():
                continue
            try:
                rec = json.loads(line)
            except ValueError as exc:
                raise ValueError(f"{path.name} satır {n}: geçerli JSON değil ({exc})") from exc
            if not isinstance(rec, dict):
                raise ValueError(f"{path.name} satır {n}: kayıt bir JSON nesnesi olmalı")
            yield rec


def _text(value: object) -> str:
    return value.strip() if isinstance(value, str) else ""


def load_term_dictionary(path: Path) -> list[TermGroup]:
    """One group per line: ``{"term", "equivalents": [...], "domain", "note"}`` (ADR-052)."""
    groups: list[TermGroup] = []
    for rec in _jsonl(path):
        term = _text(rec.get("term"))
        if not term:
            continue
        raw = rec.get("equivalents")
        items = raw if isinstance(raw, list) else []
        groups.append(
            TermGroup(
                term=term,
                equivalents=tuple(e for e in map(_text, items) if e),
                domain=_text(rec.get("domain")),
                note=_text(rec.get("note")),
            )
        )
    return groups


def load_stopword_file(path: Path) -> frozenset[str]:
    """One word per line: ``{"word", "group"}``; the group only documents why (ADR-052)."""
    if not path.exists():
        log.warning("Durak kelime dosyası bulunamadı: %s", path)
        return frozenset()
    return load_stopwords(_text(rec.get("word")) for rec in _jsonl(path))


# ---------------------------------------------------------------- word lists on the admin screen
# The term dictionary and stopwords are edited on the admin screen (ADR-052): records as
# the screen shows them, checked and written whole (atomic, previous file kept as .bak).

MAX_LIST_ITEMS = 5000


def read_terms(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [
        {"term": g.term, "equivalents": list(g.equivalents), "domain": g.domain, "note": g.note}
        for g in load_term_dictionary(path)
    ]


def read_stopwords(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    rows = []
    for rec in _jsonl(path):
        if word := _text(rec.get("word")):
            rows.append({"word": word, "group": _text(rec.get("group"))})
    return rows


def _items(items: object) -> list[dict[str, Any]]:
    if not isinstance(items, list) or not all(isinstance(i, dict) for i in items):
        raise ValueError("Liste bekleniyordu")
    if len(items) > MAX_LIST_ITEMS:
        raise ValueError(f"En fazla {MAX_LIST_ITEMS} kayıt")
    return items


def check_terms(items: object) -> list[dict[str, Any]]:
    """Clean term groups; an empty term, a group without equivalents or a term given
    twice (Turkish case-folded) is an error naming the row."""
    out: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    for n, rec in enumerate(_items(items), 1):
        term = _text(rec.get("term"))
        raw = rec.get("equivalents")
        equivalents = list(
            dict.fromkeys(e for e in map(_text, raw if isinstance(raw, list) else []) if e)
        )
        if not term:
            raise ValueError(f"Satır {n}: terim boş")
        if not equivalents:
            raise ValueError(f"“{term}”: en az bir eşdeğer gerekli")
        if (key := fold(term)) in seen:
            raise ValueError(f"“{term}” iki kez var (satır {seen[key]} ve {n})")
        seen[key] = n
        out.append(
            {
                "term": term,
                "equivalents": equivalents,
                "domain": _text(rec.get("domain")),
                "note": _text(rec.get("note")),
            }
        )
    return out


def check_stopwords(items: object) -> list[dict[str, str]]:
    """Clean stopwords: one word each (they are dropped token by token), no repeats."""
    out: list[dict[str, str]] = []
    seen: set[str] = set()
    for rec in _items(items):
        word = _text(rec.get("word"))
        if not word:
            continue
        if len(word.split()) > 1:
            raise ValueError(f"“{word}”: durak kelime tek kelime olmalı")
        if (key := fold(word)) in seen:
            continue
        seen.add(key)
        out.append({"word": word, "group": _text(rec.get("group"))})
    return out


# Each list goes out to Excel and comes back the same way: one sheet, one record per row;
# a term's equivalents share one cell, comma-separated (the screen splits them the same way).
WORD_SHEETS: dict[str, tuple[str, tuple[tuple[str, str, tuple[str, ...]], ...]]] = {
    "terms": ("Terimler", (
        ("term", "Terim", ("terim", "term")),
        ("equivalents", "Eşdeğerler", ("esdegerler", "esdeger", "equivalents")),
        ("domain", "Alan", ("alan", "domain")),
        ("note", "Not", ("not", "note", "aciklama")),
    )),
    "stopwords": ("DurakKelimeler", (
        ("word", "Kelime", ("kelime", "word", "durakkelime")),
        ("group", "Grup", ("grup", "group")),
    )),
}  # fmt: skip
_EQUIVALENT_SPLIT = re.compile(r"[,;|\n]")


def words_from_workbook(kind: str, path: Path) -> list[dict[str, Any]]:
    """A word list exported to Excel (and maybe edited) back as records for
    ``check_terms`` / ``check_stopwords``. Raises ValueError (Turkish) when it does not fit."""
    title, fields = WORD_SHEETS[kind]
    try:
        sheets = read_workbook(path, MAX_LIST_ITEMS + 1, max_cols=20)
    except TooManyRows:
        raise ValueError(f"Dosyada çok fazla satır var (en fazla {MAX_LIST_ITEMS}).") from None
    except Exception as exc:  # not an Excel file, a damaged one…
        raise ValueError("Dosya Excel (.xlsx) olarak okunamadı.") from exc
    rows = sheets.get(title) or next(iter(sheets.values()), [])
    header = {fold(_cell(h)).replace(" ", ""): i for i, h in enumerate(rows[0] if rows else [])}
    where: dict[str, int] = {}
    for key, name, aliases in fields:
        col = next((header[a] for a in aliases if a in header), None)
        if col is not None:
            where[key] = col
        elif key in ("term", "equivalents", "word"):
            raise ValueError(f"“{name}” sütunu bulunamadı; dışa aktarılan dosyanın başlıklarını "
                             f"koruyun ({', '.join(f[1] for f in fields)}).")  # fmt: skip
    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        rec: dict[str, Any] = {k: _cell(row[i]) if i < len(row) else "" for k, i in where.items()}
        if not any(rec.values()):
            continue
        if kind == "terms":
            parts = _EQUIVALENT_SPLIT.split(str(rec.get("equivalents", "")))
            rec["equivalents"] = [p.strip() for p in parts if p.strip()]
        out.append(rec)
    if not out:
        raise ValueError(f"“{title}” sayfasında kayıt yok.")
    return out


def words_to_workbook(kind: str, items: Sequence[Mapping[str, Any]], out: Path) -> Path:
    """One word list as a one-sheet workbook that ``words_from_workbook`` reads back."""
    title, fields = WORD_SHEETS[kind]
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    ws = wb.create_sheet(title)
    ws.append([name for _, name, _ in fields])
    for rec in items:
        ws.append([
            ", ".join(v) if isinstance(v := rec.get(key), list) else (v or None)
            for key, _, _ in fields
        ])  # fmt: skip
    formulas_as_text(ws)
    for row in ws.iter_rows(max_row=1):
        for cell in row:
            cell.font = Font(bold=True)
    widths = {"Terim": 28, "Eşdeğerler": 70, "Alan": 16, "Not": 50, "Kelime": 24, "Grup": 24}
    for i, (_, name, _) in enumerate(fields, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 20)
    ws.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def formulas_as_text(ws: Worksheet) -> None:
    """A text starting with "=" stays a text: openpyxl would write it as a formula."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"


def write_jsonl(path: Path, rows: Iterable[Mapping[str, Any]]) -> None:
    """Written whole or not at all; the previous file is kept as ``.bak``."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8", newline="\n") as fh:
        for r in rows:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    if path.exists():
        os.replace(path, path.with_suffix(path.suffix + ".bak"))
    os.replace(tmp, path)


# --------------------------------------------------------------------------- list search input

MAX_TERMS = 200  # one list is one sitting; a longer one is split by the user
MAX_TERM_CHARS = 500  # longer cells are pasted paragraphs, not search terms
MAX_LIST_ROWS = MAX_TERMS * 10  # blank and repeated rows included
MAX_LIST_COLUMNS = 50  # wide enough to see that a list is not one column
_HEADER_WORDS = frozenset({
    "terim", "terimler", "talep", "talepler", "soru", "sorular", "arama", "kavram", "kavramlar",
    "term", "terms", "query", "queries", "request", "requests", "liste", "aranacak",
    "aranacaklar", "ihtiyac", "konu",
})  # fmt: skip
TERM_FORMAT = "A sütununda, ilk satırda bir başlık (ör. “Terim”), altında her satırda bir terim"


@dataclass(slots=True)
class TermList:
    """A term list read by the one rule set (ADR-033): what will be searched, the header
    that was not, and everything the reader should know about what was left out."""

    terms: list[str]
    header: str
    notes: list[str]


def _column_letter(i: int) -> str:
    letters = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        letters = chr(65 + r) + letters
    return letters


def parse_term_sheet(
    rows: Sequence[Sequence[object]], other_sheets: Sequence[str] = ()
) -> TermList:
    """The first sheet of a term list -> the terms, or a ValueError that says how to fix
    the file. The format is fixed so nothing is guessed: one column (A), the first row is
    always the header and is never searched, one term per row below it."""
    cells = [[" ".join(_cell(c).split()) for c in row] for row in rows]
    filled = sorted({j for row in cells for j, c in enumerate(row) if c})
    if not filled:
        raise ValueError(f"Dosya boş. Beklenen biçim: {TERM_FORMAT}.")
    if filled != [0]:
        where = ", ".join(_column_letter(j) for j in filled)
        if 0 not in filled:
            raise ValueError(f"Terimler A sütununda olmalı; dosyada veri {where} sütununda. "
                             f"Beklenen biçim: {TERM_FORMAT}.")  # fmt: skip
        raise ValueError(f"Liste tek sütunlu olmalı; dosyada {where} sütunlarında veri var. "
                         f"Terimleri A sütununa yazıp diğer sütunları silin. Beklenen biçim: "
                         f"{TERM_FORMAT}.")  # fmt: skip
    column = [row[0] if row else "" for row in cells]
    while column and not column[-1]:
        column.pop()
    header = column[0]
    if not header:
        raise ValueError(f"İlk satır başlık olmalı; A1 hücresi boş. Beklenen biçim: {TERM_FORMAT}.")
    terms: list[str] = []
    seen: set[str] = set()
    blanks = repeats = 0
    for n, text in enumerate(column[1:], 2):
        if not text:
            blanks += 1
            continue
        if len(text) > MAX_TERM_CHARS:
            raise ValueError(f"{n}. satır çok uzun ({len(text)} karakter). Her satıra kısa bir "
                             f"terim ya da soru yazın (en fazla {MAX_TERM_CHARS}).")  # fmt: skip
        key = fold(text)
        if key in seen:
            repeats += 1
            continue
        seen.add(key)
        terms.append(text)
    if not terms:
        raise ValueError(f"Başlığın (“{header}”) altında aranacak terim yok. "
                         f"Beklenen biçim: {TERM_FORMAT}.")  # fmt: skip
    if len(terms) > MAX_TERMS:
        raise ValueError(f"Listede {len(terms)} terim var; tek seferde en fazla {MAX_TERMS}. "
                         "Listeyi bölüp ayrı ayrı yükleyin.")  # fmt: skip
    notes: list[str] = []
    if fold(header).strip(" :.") not in _HEADER_WORDS:
        notes.append(f"İlk satır (“{header}”) başlık sayıldı ve aranmadı.")
    if blanks:
        notes.append(f"{blanks} boş satır atlandı.")
    if repeats:
        notes.append(f"{repeats} tekrar eden terim atlandı (ilk geçtiği satır arandı).")
    if other_sheets:
        notes.append(f"Yalnız ilk sayfa okundu; şu sayfalardaki veriler alınmadı: "
                     f"{', '.join(other_sheets)}.")  # fmt: skip
    return TermList(terms, header, notes)


def load_term_list(path: Path) -> TermList:
    """A term list from an Excel file (first sheet; the other sheets are only checked)."""
    try:
        sheets = read_workbook(path, MAX_LIST_ROWS + 1, MAX_LIST_COLUMNS)
    except TooManyRows as exc:
        raise ValueError(f"Dosyada çok fazla satır var (en fazla {MAX_LIST_ROWS}); listeyi "
                         "bölüp ayrı ayrı yükleyin.") from exc  # fmt: skip
    except Exception as exc:  # not an Excel file, a damaged one, an old .xls…
        raise ValueError(f"Dosya Excel (.xlsx) olarak okunamadı. Beklenen biçim: "
                         f"{TERM_FORMAT}.") from exc  # fmt: skip
    if not sheets:
        raise ValueError(f"Dosyada sayfa yok. Beklenen biçim: {TERM_FORMAT}.")
    names = list(sheets)
    return parse_term_sheet(sheets[names[0]], [n for n in names[1:] if sheets[n]])
