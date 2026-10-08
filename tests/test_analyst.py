"""ADR-029: analyst flow — catalog, validation of everything the model writes, fallback.
Scripted model replies; no server needed."""

from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from vsa.config import Settings
from vsa.llm.analyst import (
    MentionChecker,
    build_answer,
    build_catalog,
    column_caveats,
    table_material,
)
from vsa.models import DictColumn, Dictionary, Verdict
from vsa.pipeline import Engine, Resources
from vsa.report.excel import write_ask_report


def col(i: int, obj: str, name: str, desc: str, schema: str = "CON") -> DictColumn:
    return DictColumn(i, "EDWDM", schema, obj, name, desc, dataset_group="Kart")


COLUMNS = [
    col(0, "vCreditCardList", "CardRefNumber", "Kartın tekil referans numarası."),
    col(1, "vCreditCardList", "CustomerId", "Ek kart sahibinin müşteri numarası."),
    col(2, "vCreditCardList", "MainCustomerId", "Asıl kart müşterisinin numarası."),
    col(3, "vCreditCardList", "CardStatusCodeName", "Kartın güncel statüsü."),
    col(4, "vChurnInput", "NumberOfCreditCards", "Müşterinin kredi kartı sayısı.", "AIS"),
    col(5, "vChurnInput", "CustomerId", "Müşteri numarası.", "AIS"),
]
LIST = "EDWDM.CON.vCreditCardList"
CHURN = "EDWDM.AIS.vChurnInput"


def objects() -> dict[str, list[DictColumn]]:
    out: dict[str, list[DictColumn]] = {}
    for c in COLUMNS:
        out.setdefault(c.object_key, []).append(c)
    return out


class ScriptedClient:
    """Answers the two analyst steps in order; None simulates a failed call."""

    available = True
    model = "fake-analyst"
    down = ""  # health(): "" = the server answers

    def health(self) -> str:
        return self.down

    def __init__(self, *replies: dict[str, Any] | None) -> None:
        self.replies = list(replies)
        self.systems: list[str] = []
        self.users: list[str] = []

    def chat_json(
        self, system: str, user: str, schema: Mapping[str, Any], *, max_tokens: int = 1024
    ) -> dict[str, Any] | None:
        self.systems.append(system)
        self.users.append(user)
        return self.replies.pop(0) if self.replies else None


def engine(client: Any) -> Engine:
    settings = Settings()
    dictionary = Dictionary(COLUMNS, "sozluk.xlsx", "abc-6")
    return Engine(dictionary, settings, Resources(), llm=client)


SHORTLIST = {
    "interpretation": "Kart bazında künye bilgisi isteniyor.",
    "candidates": [{"id": "T2", "why": "kart listesi"}, {"id": "T99", "why": "uydurma"}],
    "confusables": [{"id": "T1", "why": "model girdisi"}],
    "search_terms": ["CreditCard", "kart"],
}


def analyst_reply(**overrides: Any) -> dict[str, Any]:
    reply: dict[str, Any] = {
        "verdict": "VAR",
        "summary": "VAR. Ana kaynak EDWDM.CON.vCreditCardList.",
        "recommendations": [
            {
                "id": "T2",
                "covers": "Kart künyesi",
                "columns": ["cardrefnumber", "MainCustomerId", "CardColour"],
                "reason": "T2 tablosu kart başına bir satır tutar. Asıl müşteri MainCustomerId.",
                "caveat": "CustomerId ek kart sahibidir.",
                "usage": "Kart künyesi isteniyorsa ilk tercih.",
                "confidence": 0.9,
            }
        ],
        "design": ["Tek tablo yeterli: T2 (vCreditCardList)."],
        "attention": ["vCreditCardList.CustomerId ek kart sahibini gösterir."],
        "notes": [
            {"scope": "Uyarı", "title": "Model girdisi", "text": "vChurnInput kart listesi değil."},
            {"scope": "Top 5 dışı", "title": "Uydurma", "text": "EDWDM.CON.vCardMagic bak."},
        ],
    }
    reply.update(overrides)
    return reply


class TestCatalog:
    def test_every_table_with_all_columns(self) -> None:
        cat = build_catalog(objects())
        assert cat.ids == {"T1": CHURN, "T2": LIST}
        assert cat.keys[LIST] == "T2"
        line = cat.text.splitlines()[1]
        assert line.startswith("T2 | EDWDM.CON.vCreditCardList | Kart | 4 kolon |")
        assert "MainCustomerId" in line and "CardStatusCodeName" in line

    def test_wide_table_keeps_keys_and_relevant_descriptions(self) -> None:
        cols = objects()[LIST]
        text = table_material(LIST, "T2", cols, {3: 1.0}, 200, full_limit=2)
        assert "- CardRefNumber: Kartın tekil" in text  # structural key on a tie
        assert "- CardStatusCodeName: Kartın güncel" in text  # most relevant first
        assert "Diğer 2 kolon (yalnız ad): CustomerId, MainCustomerId" in text


class TestMentions:
    checker = MentionChecker.build(COLUMNS)

    def test_known_names_pass(self) -> None:
        text = f"{LIST} ve {LIST}.CardRefNumber ile vCreditCardList.MainCustomerId; vPRC* hariç."
        assert self.checker.unknown(text) == []

    def test_unknown_names_found(self) -> None:
        assert self.checker.unknown("EDWDM.CON.vCardMagic") == ["EDWDM.CON.vCardMagic"]
        assert self.checker.unknown(f"{LIST}.Colour") == [f"{LIST}.Colour"]
        assert self.checker.unknown("vCreditCardList.Colour") == ["vCreditCardList.Colour"]
        assert self.checker.unknown("vCardMagic tablosu") == ["vCardMagic"]

    def test_clean_removes_only_the_bad_sentence(self) -> None:
        text, dropped = self.checker.clean("İlk cümle doğru. vCardMagic uydurma. Son cümle.")
        assert (text, dropped) == ("İlk cümle doğru. Son cümle.", 1)


class TestBuildAnswer:
    def build(self, reply: dict[str, Any], min_conf: float = 0.5) -> Any:
        cat = build_catalog(objects())
        return build_answer(
            reply, {"T2": LIST}, objects(), MentionChecker.build(COLUMNS), 5, min_conf, cat.ids
        )

    def test_validated_recommendation(self) -> None:
        a = self.build(analyst_reply())
        (rec,) = a.recommendations
        assert rec.object_key == LIST
        # case-insensitive match kept, unknown column dropped
        assert [c.column for c in rec.columns] == ["CardRefNumber", "MainCustomerId"]
        assert rec.reason.startswith("vCreditCardList tablosu")  # catalog id written as name
        assert a.design == ["Tek tablo yeterli: vCreditCardList."]  # no "vX (vX)" echo
        assert a.verdict is Verdict.FOUND
        assert [n.title for n in a.notes] == ["Model girdisi"]  # made-up table's note removed
        assert a.dropped == 2  # CardColour + the sentence naming vCardMagic

    def test_names_in_place_of_ids(self) -> None:
        recs = analyst_reply()["recommendations"]
        for written in (LIST, "vCreditCardList", "`T2`"):
            a = self.build(analyst_reply(recommendations=[{**recs[0], "id": written}]))
            assert [r.object_key for r in a.recommendations] == [LIST], written

    def test_ids_outside_the_shortlist_are_refused(self) -> None:
        recs = analyst_reply()["recommendations"]
        a = self.build(analyst_reply(recommendations=[{**recs[0], "id": "T1"}]))
        assert a.recommendations == [] and a.verdict is Verdict.NOT_FOUND
        assert a.summary.startswith("BULUNAMADI.")

    def test_low_confidence_is_not_shown(self) -> None:
        recs = analyst_reply()["recommendations"]
        a = self.build(analyst_reply(recommendations=[{**recs[0], "confidence": 0.4}]))
        assert a.recommendations == []

    def test_typo_remarks_are_dropped(self) -> None:
        recs = analyst_reply()["recommendations"]
        caveat = "CustomerId ek kart sahibidir. Kolon adında 'Dept' yazım hatası var."
        a = self.build(analyst_reply(
            recommendations=[{**recs[0], "caveat": caveat}],
            attention=["Dept/Debt yazım hatasına dikkat.", "Ek kart ayrımı yapılmalı."],
        ))  # fmt: skip
        assert a.recommendations[0].caveat == "CustomerId ek kart sahibidir."
        assert a.attention == ["Ek kart ayrımı yapılmalı."]

    def test_summary_prefix_follows_verdict(self) -> None:
        a = self.build(analyst_reply(verdict="KISMEN VAR", summary="VAR — birleştirme gerekir."))
        assert a.summary == "KISMEN VAR. birleştirme gerekir."


def test_column_caveats_pick_sentences_naming_the_column() -> None:
    texts = [
        "CustomerId ek kart sahibidir. MainCustomerId asıl müşteridir.",
        "vCreditCardList.CustomerId ile birleştirme yapılmamalı.",
    ]
    assert column_caveats("CustomerId", texts) == [
        "CustomerId ek kart sahibidir.",
        "vCreditCardList.CustomerId ile birleştirme yapılmamalı.",
    ]
    assert column_caveats("MainCustomerId", texts) == ["MainCustomerId asıl müşteridir."]
    assert column_caveats("CardRefNumber", texts) == []
    # "ör." does not end the sentence
    sentence = "Oran türetilir (ör. vCreditCardLimit) CardLimit ile."
    assert column_caveats("CardLimit", [sentence]) == [sentence]


def test_family_search_adds_overlooked_tables() -> None:
    """A table the model missed in the catalog is read when its columns carry one of the
    request's information families; every candidate shows which concepts it carries."""
    cols = [*COLUMNS, col(6, "vCardTurnover", "TrnAmount", "Kart işlem tutarı (turnover).", "TRX")]
    families = [{"name": "kart işlemleri", "terms": ["Turnover", "TrnAmount"]}]
    client = ScriptedClient({**SHORTLIST, "families": families}, analyst_reply())
    settings = Settings()
    e = Engine(Dictionary(cols, "sozluk.xlsx", "abc-7"), settings, Resources(), llm=client)
    r = e.analyze("müşterinin kredi kartı bilgisi")
    step2 = client.users[1]
    assert "### T3 · EDWDM.TRX.vCardTurnover" in step2
    assert "Aile aramasıyla eklendi: kart işlemleri" in step2
    assert "Talep kavramları (kelime eşleşmesi):" in step2
    assert any("vCardTurnover (kart işlemleri)" in m for m in r.method)
    # the confusable table is never re-added as a candidate
    assert "### T1 ·" not in step2


class TestEngine:
    def test_analyst_answer(self, tmp_path: Path) -> None:
        client = ScriptedClient(SHORTLIST, analyst_reply())
        r = engine(client).analyze("müşterinin kredi kartı bilgisi")
        assert r.analyst and r.verdict is Verdict.FOUND
        assert [m.object_key for m in r.objects] == [LIST]
        m = r.objects[0]
        assert m.covers == "Kart künyesi" and m.score == 0.9 and m.llm_confidence == 0.9
        assert {h.role for h in m.columns} == {"yapısal"}
        by_name = {h.col.column: h.caveats for h in m.columns}
        # warnings name CustomerId, not the two recommended columns
        assert by_name == {"CardRefNumber": [], "MainCustomerId": []}
        assert r.design and r.attention
        assert r.notes[-1].title == "Doğrulama"
        assert r.dropped_by_validation == 3  # T99, the made-up column and table
        # step 1 sees the catalog in the system message, step 2 the full columns
        assert "T2 | EDWDM.CON.vCreditCardList" in client.systems[0]
        assert "- MainCustomerId: Asıl kart müşterisinin numarası." in client.users[1]
        assert "vChurnInput" in client.users[1]  # look-alike table as evidence

        path = write_ask_report(r, tmp_path / "r.xlsx")
        wb = load_workbook(path)
        assert wb.sheetnames == ["Özet", "Öneriler"]
        cells = [str(c.value) for row in wb["Özet"].iter_rows() for c in row if c.value]
        assert "Öneri Özeti" in cells and "Güven Skoru Ölçeği" in cells
        assert "Önerilen Kurgu" not in cells and "Yöntem Notları" not in cells
        assert "Kart künyesi" in cells

    def test_nothing_found(self) -> None:
        empty = {**SHORTLIST, "candidates": [], "confusables": []}
        r = engine(ScriptedClient(empty)).analyze("uzay gemisi yakıtı")
        assert r.analyst and r.verdict is Verdict.NOT_FOUND and r.objects == []

    def test_model_failure_says_why_and_lists_nothing(self) -> None:
        client = ScriptedClient(None)
        client.last_error = "HTTP 400: request (65019 tokens) exceeds the available context size"  # type: ignore[attr-defined]
        r = engine(client).analyze("kredi kartı")
        assert not r.analyst and r.objects == [] and r.near_misses == []  # ADR-038
        assert "bağlam penceresi" in r.fallback  # the user is told why
        assert r.summary.startswith("Analiz yapılamadı")

    def test_second_step_failure_falls_back(self) -> None:
        r = engine(ScriptedClient(SHORTLIST, None)).analyze("kredi kartı")
        assert not r.analyst and r.objects == []

    def test_unreachable_server_is_not_waited_on(self) -> None:
        """ADR-034: a dead server is found by the quick check; the model is never asked."""
        client = ScriptedClient(SHORTLIST, analyst_reply())
        client.down = "Model sunucusuna ulaşılamadı (http://spark:8000/v1): bağlantı reddedildi"
        r = engine(client).analyze("kredi kartı")
        assert not r.analyst and client.systems == [] and r.objects == []
        assert r.fallback == client.down  # said once, as the check said it

    def test_rule_answer_confidence_is_discounted(self) -> None:
        """ADR-034: the rule engine's own answer (now only measured, ADR-038) never looks
        as sure as an analyst's."""
        r = engine(ScriptedClient(None)).rule_answer("kredi kartı")
        assert not r.analyst and r.confidence_factor == 0.8 and r.objects
        for m in [*r.objects, *r.near_misses]:
            assert m.rule_score is not None and abs(m.score - round(m.rule_score * 0.8, 4)) < 1e-9
        assert any("× 0.8" in line for line in r.method)
        ok = engine(ScriptedClient(SHORTLIST, analyst_reply())).analyze("kredi kartı")
        assert ok.analyst and ok.confidence_factor == 1.0

def test_core_concept_missing_means_not_found() -> None:
    """ADR-041: look-alike tables are not an answer when the core concept is missing."""
    r = engine(ScriptedClient(SHORTLIST, analyst_reply(
        core_concept="hane geliri", core_found=False, verdict="KISMEN VAR"))
    ).analyze("müşterinin kredi kartı bilgisi")  # fmt: skip
    assert r.analyst and r.verdict is Verdict.NOT_FOUND and r.objects == []
    assert any(n.scope == "İlgili" and "hane geliri" in n.text for n in r.notes)
    said = engine(ScriptedClient(SHORTLIST, analyst_reply(verdict="BULUNAMADI"))).analyze(
        "müşterinin kredi kartı bilgisi"
    )
    assert said.verdict is Verdict.NOT_FOUND and said.objects == []  # the model's own word


def test_core_holder_leads() -> None:
    """ADR-043: the table holding the core columns goes first, shown with the top confidence."""
    from types import SimpleNamespace

    from vsa.llm.analyst import Recommendation, core_first

    def rec(key: str, conf: float) -> Recommendation:
        return Recommendation(key, [], "-", "-", "-", "-", conf)

    recs = [rec("DB.S.vSummary", 0.9), rec("DB.S.vSpending", 0.7), rec("DB.S.vOther", 0.6)]
    spending, other = (SimpleNamespace(object_key=f"DB.S.{n}") for n in ("vSpending", "vOther"))
    core = [spending, spending, other]
    ranked = core_first(recs, core)  # type: ignore[arg-type]
    assert [r.object_key for r in ranked] == ["DB.S.vSpending", "DB.S.vOther", "DB.S.vSummary"]
    assert ranked[0].confidence == 0.9 and ranked[1].confidence == 0.6
    assert core_first(recs, []) == recs  # no core columns: the model's order stands
    # a tie on core columns: the general table before the narrow-population one
    tie = [rec("DB.S.vPrivateBanking", 0.9), rec("DB.S.vBalanceDaily", 0.8)]
    both = [SimpleNamespace(object_key=k) for k in ("DB.S.vPrivateBanking", "DB.S.vBalanceDaily")]
    first = core_first(tie, both, {"DB.S.vPrivateBanking"})  # type: ignore[arg-type]
    assert [r.object_key for r in first] == ["DB.S.vBalanceDaily", "DB.S.vPrivateBanking"]
    alone = core_first(tie, both[:1], {"DB.S.vPrivateBanking"})  # type: ignore[arg-type]
    assert alone[0].object_key == "DB.S.vPrivateBanking"  # it alone holds the core: it leads
    outside = [SimpleNamespace(object_key="DB.S.vElsewhere")]
    assert core_first(recs, outside) == recs  # type: ignore[arg-type]


def test_core_concept_needs_a_real_column() -> None:
    """ADR-042: the core concept must be carried by a named candidate column."""
    named = engine(ScriptedClient(SHORTLIST, analyst_reply(
        core_concept="asıl kart sahibi", core_columns=["T2.MainCustomerId"], core_found=True))
    ).analyze("müşterinin kredi kartı bilgisi")  # fmt: skip
    assert named.verdict is Verdict.FOUND and named.objects
    for bad in ([], ["T2.HaneGeliri"], ["T99.MainCustomerId"]):
        r = engine(ScriptedClient(SHORTLIST, analyst_reply(
            core_concept="hane geliri", core_columns=bad, core_found=True))
        ).analyze("müşterinin kredi kartı bilgisi")  # fmt: skip
        assert r.verdict is Verdict.NOT_FOUND and r.objects == [], bad
        assert any(n.scope == "İlgili" and "hane geliri" in n.text for n in r.notes)



def test_partial_answer_names_what_is_missing() -> None:
    """A KISMEN VAR answer carries the missing parts to show by the verdict; VAR none."""
    partial = engine(ScriptedClient(SHORTLIST, analyst_reply(
        verdict="KISMEN VAR", missing=["kart rengi geçmişi", ""], core_found=True))
    ).analyze("müşterinin kredi kartı bilgisi")  # fmt: skip
    assert partial.verdict is Verdict.PARTIAL and partial.missing == ["kart rengi geçmişi"]
    full = engine(ScriptedClient(SHORTLIST, analyst_reply(missing=["gereksiz"]))).analyze(
        "müşterinin kredi kartı bilgisi"
    )
    assert full.verdict is Verdict.FOUND and full.missing == []
