"""End-to-end: M1 acceptance on the golden set, report and CLI (needs the real dictionary)."""

from __future__ import annotations

import logging
from pathlib import Path

import pytest
from openpyxl import load_workbook

from vsa.config import Settings
from vsa.evaluation import evaluate, load_golden
from vsa.models import Verdict
from vsa.pipeline import Engine
from vsa.report.excel import report_path, slugify, write_ask_report

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def engine() -> Engine:
    from .conftest import REAL_DICTIONARY

    if not REAL_DICTIONARY.exists():
        pytest.skip("Gerçek sözlük yok (data/ gitignore'da)")
    logging.disable(logging.WARNING)
    settings = Settings()
    settings.dictionary.path = str(REAL_DICTIONARY)
    return Engine.from_dictionary_file(settings)


def test_m1_acceptance(engine: Engine) -> None:
    """HANDOVER §17 M1: expected object in the top 5 for every golden ask item;
    no trap violations; every negative item answered BULUNAMADI (ADR-006).
    Items taken from the hand-made chat analyses measure the analyst flow (ADR-029),
    which needs a model; the LLM-free core is held to the M1 items."""
    golden = [
        item
        for item in load_golden(ROOT / "tests" / "golden_set.yaml")
        if item.get("source") != "chat-analiz-2026-09-26"
    ]
    report = evaluate(engine, golden, load_golden(ROOT / "tests" / "negative_set.yaml"))
    asks = report.group("ask")
    assert asks
    misses = [i.id for i in asks if not i.rank or i.rank > 5]
    assert misses == [], misses
    assert report.trap_violations == 0
    assert report.false_answer_rate == 0.0


def test_ask_report(engine: Engine, tmp_path: Path) -> None:
    result = engine.rule_answer("Kredi kartı limit doluluk oranı", top_n=3)
    assert result.verdict is not Verdict.NOT_FOUND
    assert 0 < len(result.objects) <= 3
    assert result.dropped_by_validation == 0
    keys = engine.column_keys
    assert all(h.col.key in keys for m in result.objects for h in m.columns)

    path = write_ask_report(result, report_path(tmp_path, "ask", result.query))
    wb = load_workbook(path)
    assert wb.sheetnames == ["Özet", "Öneriler"]
    ws = wb["Öneriler"]
    assert ws["A3"].value == "Sıra"
    assert ws["A3"].font.name == "Arial"
    assert ws["I4"].number_format == "0.00"
    assert ws.freeze_panes == "A4"
    assert str(wb["Özet"]["A2"].value).startswith("Kaynak sözlük:")
    assert str(wb["Özet"]["A3"].value).startswith("Cevabı üreten: Kural motoru")  # ADR-034
    said = {row[0].value: row[1].value for row in wb["Özet"].iter_rows(min_row=5, max_row=7)}
    assert said["Talep"] == result.query and said["Sonuç"] == result.summary


def test_irrelevant_query_does_not_crash(engine: Engine) -> None:
    result = engine.analyze("zzzz qqqq")
    assert result.verdict is Verdict.NOT_FOUND
    assert result.objects == []


def test_slug() -> None:
    assert slugify("Kredi kartı limit doluluk oranı?") == "kredi_karti_limit_doluluk_orani"


def test_list_report(engine: Engine, tmp_path: Path) -> None:
    """A term list, each term through the question flow, one combined report (ADR-033)."""
    from vsa.models import ListItem, ListResult
    from vsa.report.excel import write_list_report

    terms = ["kredi kartı limit doluluk oranı", "uzay gemisi yakıt seviyesi"]
    items = [ListItem(i, t, engine.rule_answer(t)) for i, t in enumerate(terms, 1)]
    items.append(ListItem(3, "durdurulan terim"))
    result = ListResult("liste.xlsx", items, "sozluk.xlsx", "v1", "2026-10-01 10:00",
                        cancelled=True)  # fmt: skip
    wb = load_workbook(write_list_report(result, tmp_path / "l.xlsx"))
    assert wb.sheetnames == ["Özet", "Öneriler"]
    summary = [[c.value for c in row] for row in wb["Özet"].iter_rows()]
    states = {r[1]: r[2] for r in summary if r and r[0] in (1, 2, 3)}
    assert states["uzay gemisi yakıt seviyesi"] == "BULUNAMADI"
    assert states["durdurulan terim"] == "DURDURULDU"
    assert states["kredi kartı limit doluluk oranı"] in ("VAR", "KISMEN VAR")
    first = next(r for r in summary if r and r[0] == 1)
    link = wb["Özet"].cell(row=summary.index(first) + 1, column=2).hyperlink
    assert link is not None and "Öneriler" in str(link.location or link.target)
    rows = [[c.value for c in row] for row in wb["Öneriler"].iter_rows(min_row=2)]
    assert {r[0] for r in rows} == {1, 2, 3}  # every term has at least one row
    assert all(r[1] for r in rows)  # and carries its term
