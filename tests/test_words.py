"""Term dictionary and stopwords edited on the admin screen (ADR-052)."""

from __future__ import annotations

import io
import json
import shutil
import threading
import urllib.error
import urllib.request
from collections.abc import Iterator
from http.server import ThreadingHTTPServer
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from vsa.config import load_settings
from vsa.loader import check_stopwords, check_terms, load_term_dictionary
from vsa.pipeline import Engine
from vsa.web.server import App, make_handler

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def server(sample_store_path: Path, tmp_path: Path) -> Iterator[tuple[str, App]]:
    for name in ("terms.jsonl", "stopwords.jsonl"):
        shutil.copy(ROOT / "data" / name, tmp_path / name)
    settings = tmp_path / "settings.yaml"
    settings.write_text(
        f"dictionary:\n  store: {sample_store_path.as_posix()}\n"
        f"index:\n  dir: {(tmp_path / 'index').as_posix()}\n"
        f"expansion:\n  term_dictionary: {(tmp_path / 'terms.jsonl').as_posix()}\n"
        f"  stopwords: {(tmp_path / 'stopwords.jsonl').as_posix()}\n",
        encoding="utf-8",
    )
    s = load_settings(settings)
    Engine.from_dictionary_file(s).save()
    app = App(Engine.from_index(s), tmp_path / "out", tmp_path / "feedback.jsonl", settings)
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(app))  # no password: loopback
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}", app
    httpd.shutdown()
    httpd.server_close()


def call(url: str, body: Any = None) -> tuple[int, dict[str, Any]]:
    data = None if body is None else json.dumps(body).encode()
    req = urllib.request.Request(url, data=data, headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, json.loads(r.read())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read())


def test_lists_are_read_and_replaced(server: tuple[str, App], tmp_path: Path) -> None:
    url, app = server
    status, state = call(url + "/api/admin/words")
    assert status == 200
    terms, stops = state["terms"]["items"], state["stopwords"]["items"]
    assert len(terms) >= 45 and {"word": "ve", "group": "Genel"} in stops

    new = {"term": "zırhlı araç", "equivalents": ["panzer", " panzer "], "domain": "test"}
    terms.insert(0, new)
    status, r = call(url + "/api/admin/words/terms", {"items": terms})
    assert status == 200 and r["count"] == len(terms)
    saved = load_term_dictionary(tmp_path / "terms.jsonl")
    assert saved[0].term == "zırhlı araç" and saved[0].equivalents == ("panzer",)
    assert (tmp_path / "terms.jsonl.bak").is_file()  # previous list kept
    # the engine was rebuilt with the new group: one word asks for the other
    assert any(g.term == "zırhlı araç" for g in app.engine.resources.term_groups)

    status, r = call(url + "/api/admin/words/stopwords", {"items": [*stops, {"word": "panzer"}]})
    assert status == 200 and "panzer" in app.engine.resources.stopwords


def test_bad_lists_are_refused(server: tuple[str, App], tmp_path: Path) -> None:
    url, _ = server
    before = (tmp_path / "terms.jsonl").read_bytes()
    twice = [{"term": "Kredi", "equivalents": ["a"]}, {"term": "kredi", "equivalents": ["b"]}]
    status, r = call(url + "/api/admin/words/terms", {"items": twice})
    assert status == 400 and "iki kez" in r["error"]
    status, r = call(url + "/api/admin/words/terms", {"items": [{"term": "x"}]})
    assert status == 400 and "eşdeğer" in r["error"]
    status, r = call(url + "/api/admin/words/stopwords", {"items": [{"word": "iki kelime"}]})
    assert status == 400 and "tek kelime" in r["error"]
    status, r = call(url + "/api/admin/words/terms", {"items": "değil"})
    assert status == 400
    assert (tmp_path / "terms.jsonl").read_bytes() == before  # nothing written


def test_checks() -> None:
    assert check_stopwords([{"word": "Ve"}, {"word": "ve"}, {"word": " "}]) == [
        {"word": "Ve", "group": ""}
    ]
    (row,) = check_terms([{"term": " a ", "equivalents": ["b", "b", ""], "note": 3}])
    assert row == {"term": "a", "equivalents": ["b"], "domain": "", "note": ""}


def fetch(url: str) -> bytes:
    with urllib.request.urlopen(url, timeout=60) as r:
        assert r.headers["Content-Type"].startswith("application/vnd.openxmlformats")
        return bytes(r.read())


def upload(url: str, data: bytes) -> tuple[int, dict[str, Any]]:
    req = urllib.request.Request(url, data=data, headers={"X-Filename": "liste.xlsx"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, json.loads(r.read())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read())


def test_lists_go_to_excel_and_come_back(server: tuple[str, App], tmp_path: Path) -> None:
    url, app = server
    for kind in ("terms", "stopwords"):
        # what a save of the current list keeps (the file may hold folded repeats: mi / mı)
        check = check_terms if kind == "terms" else check_stopwords
        before = check(app.words_state()[kind]["items"])
        base = f"{url}/api/admin/words/{kind}"
        status, r = upload(base + "/import", fetch(base + "/export"))
        assert status == 200 and r["state"][kind]["items"] == before  # round trip is lossless

    # the exported sheet edited in Excel: a row added, equivalents typed with ; and ,
    book = load_workbook(io.BytesIO(fetch(url + "/api/admin/words/terms/export")))
    ws = book["Terimler"]
    ws.append(["zırhlı araç", "panzer; tank, =zırh", "test", None])
    buf = io.BytesIO()
    book.save(buf)
    status, r = upload(url + "/api/admin/words/terms/import", buf.getvalue())
    assert status == 200
    added = next(t for t in r["state"]["terms"]["items"] if t["term"] == "zırhlı araç")
    assert added["equivalents"] == ["panzer", "tank", "=zırh"]
    assert any(g.term == "zırhlı araç" for g in app.engine.resources.term_groups)


def test_bad_workbooks_are_refused(server: tuple[str, App], tmp_path: Path) -> None:
    url, _ = server
    before = (tmp_path / "stopwords.jsonl").read_bytes()
    status, r = upload(url + "/api/admin/words/stopwords/import", b"not excel")
    assert status == 400 and "okunamadı" in r["error"]
    book = Workbook()
    book.active.append(["Başka", "Sütun"])
    book.active.append(["x", "y"])
    buf = io.BytesIO()
    book.save(buf)
    status, r = upload(url + "/api/admin/words/stopwords/import", buf.getvalue())
    assert status == 400 and "Kelime" in r["error"]
    status, _ = upload(url + "/api/admin/words/other/import", buf.getvalue())
    assert status == 404
    assert (tmp_path / "stopwords.jsonl").read_bytes() == before  # nothing written
