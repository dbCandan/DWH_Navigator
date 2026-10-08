"""ADR-049/050: the app's one dictionary (jsonl) — import, read, export; Excel only for transfer."""

from __future__ import annotations

import json
import threading
import urllib.error
import urllib.request
from collections.abc import Iterator
from http.server import ThreadingHTTPServer
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from tests.conftest import write_xlsx
from vsa import dictionary_store as ds
from vsa.cli import main as cli_main
from vsa.config import Settings, load_settings
from vsa.loader import build_columns, object_profiles, read_workbook, records
from vsa.pipeline import Engine
from vsa.web.server import App, make_handler


def parse_workbook(path: Path) -> list[Any]:
    """The workbook as the template parser reads it — what an import must keep."""
    sheets = read_workbook(path)
    profiles = object_profiles(records(sheets.get("Objeler", [])))
    columns, _ = build_columns(list(records(sheets["Kolonlar"])),
                               {k: p.group for k, p in profiles.items() if p.group})  # fmt: skip
    return columns


def test_import_keeps_the_workbook(sample_dictionary_path: Path, tmp_path: Path) -> None:
    store = ds.from_workbook(sample_dictionary_path)
    assert store.meta["columns"] == 3 and store.meta["objects"] == 2
    assert any("Tekrarlanan" in w for w in store.warnings)
    path = tmp_path / "dictionary.jsonl"
    ds.save(store, path)
    d = ds.load_dictionary(path)
    assert d.columns == parse_workbook(sample_dictionary_path)
    assert d.version == store.meta["version"] and d.version.endswith("-4")
    lines = path.read_text(encoding="utf-8").splitlines()
    assert json.loads(lines[0])["type"] == "meta" and len(lines) == 1 + 2 + 4
    assert "DatabaseName" in json.loads(lines[-1])  # "DAtabaseName" stored under its template name


def test_export_round_trip_and_template(sample_store_path: Path, tmp_path: Path) -> None:
    before = ds.load_dictionary(sample_store_path)
    out = ds.export(ds.read(sample_store_path), tmp_path / "out.xlsx")
    assert parse_workbook(out) == before.columns
    ds.save(ds.from_workbook(out, "out.xlsx"), sample_store_path)  # same records: version kept
    after = ds.load_dictionary(sample_store_path)
    assert (after.columns, after.version) == (before.columns, before.version)
    assert sample_store_path.with_suffix(".jsonl.bak").is_file()

    wb = load_workbook(ds.export(None, tmp_path / "t.xlsx"))
    assert wb.sheetnames == ["Objeler", "Kolonlar"]
    assert [c.value for c in wb["Kolonlar"][1]] == list(ds.COLUMN_FIELDS)


def test_hand_edited_store_gets_its_own_version(sample_store_path: Path) -> None:
    before = ds.load_dictionary(sample_store_path).version
    text = sample_store_path.read_text(encoding="utf-8")
    sample_store_path.write_text(text.replace("Tekrar.", "Değişti."), encoding="utf-8")
    assert ds.load_dictionary(sample_store_path).version not in (before, "")


@pytest.mark.parametrize(
    ("sheets", "message"),
    [
        ({"Sayfa1": [{"a": 1}]}, "Kolonlar"),
        ({"Kolonlar": [{"DatabaseName": None, "SchemaName": "S", "ObjectName": "T",
                        "ColumnName": "C", "ColumnDescription": "d"}]}, "Geçerli kolon"),
        ({"Kolonlar": [{"SchemaName": "S"}]}, "zorunlu kolon"),
    ],
)  # fmt: skip
def test_bad_workbooks_are_refused(
    sheets: dict[str, list[dict[str, Any]]], message: str, tmp_path: Path
) -> None:
    with pytest.raises(ValueError, match=message):
        ds.from_workbook(write_xlsx(tmp_path / "x.xlsx", sheets))


def test_engine_reads_only_the_store(sample_store_path: Path) -> None:
    s = Settings()
    s.dictionary.store = str(sample_store_path)
    engine = Engine.from_dictionary_file(s)
    assert engine.dictionary.source_path.endswith("dictionary.jsonl")
    assert len(engine.dictionary.columns) == 3


def test_cli_import_and_export(sample_dictionary_path: Path, tmp_path: Path) -> None:
    settings = tmp_path / "settings.yaml"
    settings.write_text(
        f"dictionary:\n  store: {(tmp_path / 'd.jsonl').as_posix()}\n"
        f"index:\n  dir: {(tmp_path / 'index').as_posix()}\n",
        encoding="utf-8",
    )
    assert cli_main(["dictionary", "import", str(sample_dictionary_path),
                     "--settings", str(settings)]) == 0  # fmt: skip
    meta = json.loads((tmp_path / "index" / "meta.json").read_text(encoding="utf-8"))
    assert meta["columns"] == 3  # the import is the dictionary: indexed at once
    assert cli_main(["dictionary", "export", str(tmp_path / "e.xlsx"),
                     "--settings", str(settings)]) == 0  # fmt: skip
    assert (tmp_path / "e.xlsx").is_file()


# --------------------------------------------------------------------------- admin API


@pytest.fixture
def server(tmp_path: Path) -> Iterator[str]:
    """No dictionary imported yet: the app starts empty (ADR-050)."""
    settings = tmp_path / "settings.yaml"
    settings.write_text(
        f"dictionary:\n  store: {(tmp_path / 'dictionary.jsonl').as_posix()}\n"
        f"index:\n  dir: {(tmp_path / 'index').as_posix()}\n",
        encoding="utf-8",
    )
    s = load_settings(settings)
    Engine.from_dictionary_file(s).save()
    app = App(Engine.from_index(s), tmp_path / "out", tmp_path / "feedback.jsonl", settings)
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(app))  # no password: loopback
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}"
    httpd.shutdown()
    httpd.server_close()


def call(url: str, data: bytes | None = None, headers: dict[str, str] | None = None
         ) -> tuple[int, bytes]:  # fmt: skip
    req = urllib.request.Request(url, data=data, headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def test_admin_dictionary_tab(server: str, sample_dictionary_path: Path, tmp_path: Path) -> None:
    base = server
    status, body = call(base + "/api/admin/dictionary")
    state = json.loads(body)
    assert status == 200 and state["summary"] is None and not state["store"]["exists"]
    assert call(base + "/api/admin/dictionary/export")[0] == 404  # nothing imported yet
    status, body = call(base + "/api/status")
    assert json.loads(body)["columns"] == 0  # started without a dictionary

    bad = write_xlsx(tmp_path / "bad.xlsx", {"Sayfa1": [{"a": 1}]}).read_bytes()
    status, body = call(base + "/api/admin/dictionary/import", bad)
    assert status == 400 and "Kolonlar" in json.loads(body)["error"]

    status, body = call(base + "/api/admin/dictionary/import", sample_dictionary_path.read_bytes(),
                        {"X-Filename": "VeriSozlugu.xlsx"})  # fmt: skip
    r = json.loads(body)
    assert status == 200 and r["state"]["summary"]["meta"]["source"] == "VeriSozlugu.xlsx"
    status, body = call(base + "/api/status")
    assert json.loads(body)["dictionary"] == "dictionary.jsonl"
    assert json.loads(body)["columns"] == 3  # in use at once

    status, body = call(base + "/api/admin/dictionary/table/EDWDM.CMP.vCardLimitFullness")
    table = json.loads(body)
    assert status == 200 and [c["ColumnName"] for c in table["columns"]] == [
        "CardLimitFullnessToday", "CustomerPartyId", "CardLimitFullnessToday"]  # fmt: skip
    assert table["object"]["DatasetGroup"] == "Kart"
    assert call(base + "/api/admin/dictionary/table/yok.yok.yok")[0] == 404

    status, xlsx = call(base + "/api/admin/dictionary/export")
    assert status == 200 and xlsx[:2] == b"PK"
    assert call(base + "/api/admin/dictionary/template")[0] == 200
    assert call(base + "/api/admin/dictionary/source", b'{"source": "excel"}')[0] == 404  # gone
